from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = "documents/Lightweight-University-CMS-Report.docx"


def set_para_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
      run.font.name = "Aptos Display"
      run.font.color.rgb = RGBColor(127, 29, 29) if level == 1 else RGBColor(15, 23, 42)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def callout(doc, title, body):
    p = doc.add_paragraph()
    set_para_shading(p, "FEE2E2")
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor(127, 29, 29)
    p.add_run(body)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(6)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("LIGHTWEIGHT UNIVERSITY CMS PLATFORM")
r.bold = True
r.font.size = Pt(22)
r.font.color.rgb = RGBColor(127, 29, 29)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("Technical Report for Faculty Website Modernization")
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(51, 65, 85)

doc.add_paragraph()
callout(
    doc,
    "Project summary",
    "A WordPress replacement CMS for a university faculty website, built with Next.js, TypeScript, Prisma, PostgreSQL, Tiptap, Zustand, and role-based access control.",
)

heading(doc, "1. Project Overview")
doc.add_paragraph(
    "The platform provides a public faculty website and an authenticated content management dashboard. "
    "The public side supports news, announcements, admissions, programs, faculty information, about, contact, and search pages. "
    "The admin side lets non-technical staff manage posts, media, categories, navigation menus, users, draft workflow, and SEO metadata."
)

heading(doc, "2. Main Objectives")
for item in [
    "Replace WordPress dependency with a lightweight modular monolith.",
    "Improve maintainability through a feature-based Next.js App Router structure.",
    "Provide a staff-friendly dashboard for draft, edit, publish, image upload, and SEO management.",
    "Render public pages with SEO-first metadata, sitemap, and robots.txt support.",
    "Prepare deployment for Vercel with Neon or Supabase PostgreSQL.",
]:
    bullet(doc, item)

heading(doc, "3. Architecture")
doc.add_paragraph("The solution is organized into five practical layers:")
for item in [
    "Public UI: server-rendered pages for students, parents, and faculty visitors.",
    "Admin UI: React dashboard with Zustand state and a Tiptap rich text editor.",
    "API layer: Next.js route handlers with Zod validation and RBAC authorization.",
    "Data layer: Prisma ORM with PostgreSQL tables for users, posts, categories, media, and menus.",
    "Storage layer: Vercel Blob in production and local public upload fallback for development.",
]:
    bullet(doc, item)

heading(doc, "4. Database Design")
for item in [
    "users: stores account profile, email, password hash, role, avatar, active status, and timestamps.",
    "posts: stores title, slug, HTML content, excerpt, thumbnail, status, SEO fields, author, category, and publish timestamp.",
    "categories: stores editorial taxonomy with unique names and slugs.",
    "media: stores uploaded image URL, storage key, alt text, uploader, and timestamp.",
    "menus: stores dynamic public navigation labels, paths, and display order.",
]:
    bullet(doc, item)

heading(doc, "5. Functional Scope")
for item in [
    "Public website: home, news, news detail, admissions, programs, faculty, about, contact, and search.",
    "Admin dashboard: total posts, users, latest articles, draft statistics, and category overview.",
    "Post management: create, edit, delete, save draft, publish, slug, thumbnail, and SEO metadata.",
    "Media management: upload, preview, and reuse images inside article forms.",
    "User management: create Admin or Staff accounts with password confirmation.",
    "Language support: public homepage and shared navigation can switch between English and Vietnamese.",
]:
    bullet(doc, item)

heading(doc, "6. Security")
for item in [
    "Passwords are hashed with bcrypt before storage.",
    "Sessions use HTTP-only cookies with signed JWT claims.",
    "Auth.js/NextAuth credentials route is included for compatibility and future extension.",
    "API mutations validate inputs with Zod and enforce ADMIN/STAFF role permissions.",
    "Same-origin checks and simple rate limits are applied to sensitive endpoints.",
]:
    bullet(doc, item)

heading(doc, "7. UI Improvements Completed")
for item in [
    "Broken external seed images were replaced with local CMS assets.",
    "Homepage information section was redesigned for better desktop and mobile readability.",
    "Admin post editor layout was rebuilt to avoid narrow-screen overlap.",
    "Editor toolbar now resembles a compact Word-style ribbon with bold, italic, heading, list, and image insert controls.",
    "Article form now supports thumbnail upload and preview directly from the editor modal.",
]:
    bullet(doc, item)

heading(doc, "8. Deployment Plan")
for item in [
    "Local development: copy .env.example, install packages, run Prisma migration/seed, then npm run dev.",
    "Production web: deploy to Vercel using the existing build command.",
    "Production database: use Neon or Supabase PostgreSQL and set DATABASE_URL.",
    "Production storage: enable Vercel Blob with BLOB_READ_WRITE_TOKEN.",
    "Required secrets: AUTH_SECRET, JWT_SECRET, NEXT_PUBLIC_SITE_URL, and database connection string.",
]:
    bullet(doc, item)

heading(doc, "9. Verification Summary")
for item in [
    "npm run build completed successfully.",
    "npm run lint completed with only non-blocking dynamic image warnings.",
    "Public site and admin dashboard were checked locally.",
    "Demo URLs: http://localhost:3001 and http://localhost:3001/admin.",
    "Demo accounts: admin@scse.local / Admin@12345 and staff@scse.local / Staff@12345.",
]:
    bullet(doc, item)

doc.add_paragraph()
closing = doc.add_paragraph()
closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = closing.add_run("Prepared for academic review and project submission.")
r.italic = True
r.font.color.rgb = RGBColor(71, 85, 105)

doc.save(OUT)
print(OUT)
